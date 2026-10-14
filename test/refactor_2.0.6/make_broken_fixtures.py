#!/usr/bin/env python3
"""
Creates broken-input fixtures from the valid old_triplicate fixtures (copies with edited
cells). Plate files: 2.0.6 Step 1 (ID3 cases). Protocols: the concentration checks of
Step 2 / Step 5. TEMPORARY tooling for the 2.0.6 refactor (deleted before the merge into
main); it reads and writes real lab data in test/data/ (local only, git-ignored).
Run from the repository root:

    uv run python test/refactor_2.0.6/make_broken_fixtures.py   # writes into test/data/layout_2.0.6/broken/

Files provided by the user (kept as they are, not regenerated here):
    plate with descending ID3.xlsx        ID3: 4-1
    plate too many ID3.xlsx               ID3: 1-12 (old triplicate protocol: 8 IDs, 4 conditions)
    second_ligand_with diff counts.xlsm   ligand lists with 7 and 4 unique concentrations
    old layout_non descending list.xlsm   -8, -8, -9 ... -13: NOT increasing, so it passes the
                                          descending check (design 1.2); it has 6 unique values
                                          for 7 stimulations -> caught by the Step-5 unique-count check

Files created here:
    protocol_not_descending.xlsm          ligand list -8, -7, -9, -10, -11, -12, -13 (increases at row B)
    plate_ID3_non_numeric.xlsx            ID3: 1a,2
    plate_ID3_13.xlsx                     ID3: 13 (not in the protocol)
    plate_ID3_empty_entry.xlsx            ID3: 1,,3,4 (partially used plate, legitimate)
    plate_no_ID3.xlsx                     'ID3:' line removed
"""
import os

import openpyxl

REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, os.pardir))
FIXTURES = os.path.join(REPO_ROOT, "test", "data", "layout_2.0.6")
BROKEN_DIR = os.path.join(FIXTURES, "broken")  # output folder (data, git-ignored)
OLD_FOLDER = os.path.join(FIXTURES, "old_triplicate", "260611_old layout_CXCR4 trafficking n=1 CON, dQ")
PLATE_SRC = os.path.join(OLD_FOLDER, "260611 Con CXCR4-nLuc, CAAX, FYVE, Rab7, Rab11a  P=1 N=1 (MK).xlsx")
PROTOCOL_SRC = os.path.join(OLD_FOLDER, "Procotol_CXCR4_n=1.xlsm")

CONC_MARKER = "final concentration in well (log(M))"


def _id3_cell(ws):
    for row in ws.iter_rows(min_col=1, max_col=1):
        cell = row[0]
        if isinstance(cell.value, str) and cell.value.replace(" ", "").startswith("ID3:"):
            return cell
    raise RuntimeError("no ID3 line in source plate file")


def make_plate(target: str, id3_text: str | None):
    wb = openpyxl.load_workbook(PLATE_SRC)
    cell = _id3_cell(wb["Table All Cycles"])
    cell.value = None if id3_text is None else f"ID3: {id3_text}"
    wb.save(os.path.join(BROKEN_DIR, target))


def make_protocol_not_descending(target: str, values: list[float]):
    wb = openpyxl.load_workbook(PROTOCOL_SRC, keep_vba=True)
    ws = wb["Protocol"]
    marker = next(c for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str) and CONC_MARKER in c.value)  # first list = ligand 1
    for i, v in enumerate(values, start=1):
        ws.cell(row=marker.row + i, column=marker.column).value = v
    wb.save(os.path.join(BROKEN_DIR, target))


def main():
    if not (os.path.isfile(PLATE_SRC) and os.path.isfile(PROTOCOL_SRC)):
        raise SystemExit(f"Source fixtures not found in {OLD_FOLDER} - the real lab data in "
                         f"test/data/ is local only (not in the public repository).")
    os.makedirs(BROKEN_DIR, exist_ok=True)
    make_protocol_not_descending("protocol_not_descending.xlsm", [-8, -7, -9, -10, -11, -12, -13])
    make_plate("plate_ID3_non_numeric.xlsx", "1a,2")
    make_plate("plate_ID3_13.xlsx", "13")
    make_plate("plate_ID3_empty_entry.xlsx", "1,,3,4")
    make_plate("plate_no_ID3.xlsx", None)
    print("broken fixtures written to", BROKEN_DIR)


if __name__ == "__main__":
    main()
